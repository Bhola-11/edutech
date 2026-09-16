import csv
from io import BytesIO
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def export_queryset_to_csv(queryset, field_names, header_titles=None, filename='export.csv'):
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename=\"{filename}\"'
    writer = csv.writer(response)

    if header_titles:
        writer.writerow(header_titles)
    else:
        writer.writerow(field_names)

    for obj in queryset:
        row = []
        for field in field_names:
            val = obj
            for part in field.split('__'):
                val = getattr(val, part, None)
                if callable(val):
                    val = val()
            if val is None:
                val = ''
            row.append(str(val))
        writer.writerow(row)

    return response

def export_queryset_to_excel(queryset, field_names, header_titles=None, filename='export.xlsx', sheet_title='Data Export'):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    headers = header_titles if header_titles else [f.replace('__', ' ').replace('_', ' ').title() for f in field_names]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    row_idx = 2
    for obj in queryset:
        row = []
        for field in field_names:
            val = obj
            for part in field.split('__'):
                val = getattr(val, part, None)
                if callable(val):
                    val = val()
            if val is None:
                val = ''
            row.append(val)
        ws.append(row)

        for col_num in range(1, len(row) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')
        row_idx += 1

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=\"{filename}\"'
    return response
